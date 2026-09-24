"""
fix_q20_in_excel.py

For every row where
    '20. Has the child ever dropped out of school?' == 'No'
rewrite it to 'Yes, temporarily' and fill
    'If yes, what was the main reason?'
with a random plausible reason.

Writes to a NEW file (synthetic_responses_fixed.xlsx) so the original is
untouched. Re-run the submitter with --file synthetic_responses_fixed.xlsx
after inspecting the diff.
"""

from openpyxl import load_workbook
import random

# ---- config ----------------------------------------------------------
SRC       = "synthetic_responses.xlsx"
DST       = "synthetic_responses_fixed.xlsx"
SHEET     = "Responses"

Q20_COL    = "20. Has the child ever dropped out of school?"
FOLLOW_COL = "If yes, what was the main reason?"
AGE_COL    = "17. How old is the child?"

# Reasons available in the form, split by plausibility for the child's age.
REASONS_YOUNG = [   # for children 5–9 and 10–13
    "Financial constraints",
    "Need to work",
    "Household responsibilities",
    "Lack of interest",
    "Poor academic performance",
    "Distance/transportation",
    "Health reasons",
]
REASONS_OLDER = REASONS_YOUNG + ["Marriage", "Pregnancy"]  # 14–17 only

NEW_Q20_VALUE = "Yes, temporarily"   # or "Yes, permanently" — see notes

SEED = 42   # reproducible; change or remove for variety
# ---------------------------------------------------------------------

random.seed(SEED)

wb = load_workbook(SRC)
ws = wb[SHEET]

# Locate column indices from the header row.
headers = {}
for cell in ws[1]:
    if cell.value is not None:
        headers[str(cell.value)] = cell.column

for name in (Q20_COL, FOLLOW_COL, AGE_COL):
    if name not in headers:
        raise SystemExit(f"Column not found in header row: {name!r}")

q20_c     = headers[Q20_COL]
follow_c  = headers[FOLLOW_COL]
age_c     = headers[AGE_COL]

def cell_str(row, col):
    v = ws.cell(row=row, column=col).value
    return "" if v is None else str(v).strip()

changed = 0
skipped = 0

for r in range(2, ws.max_row + 1):
    q20    = cell_str(r, q20_c)
    follow = cell_str(r, follow_c)
    age    = cell_str(r, age_c)

    if not q20 and not follow and not age:
        continue  # blank trailing row

    if q20 != "No":
        if q20 in ("Yes, temporarily", "Yes, permanently") and not follow:
            # Already Yes but follow-up empty — fill it.
            pool = REASONS_OLDER if age == "14–17 years" else REASONS_YOUNG
            ws.cell(row=r, column=follow_c, value=random.choice(pool))
            changed += 1
        else:
            skipped += 1
        continue

    # q20 == "No": rewrite it so the follow-up becomes applicable.
    pool = REASONS_OLDER if age == "14–17 years" else REASONS_YOUNG
    ws.cell(row=r, column=q20_c,    value=NEW_Q20_VALUE)
    ws.cell(row=r, column=follow_c, value=random.choice(pool))
    changed += 1

wb.save(DST)
print(f"Wrote {DST}")
print(f"  rows changed : {changed}")
print(f"  rows left as-is (already consistent or non-No): {skipped}")